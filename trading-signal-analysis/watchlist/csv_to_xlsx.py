# -*- coding: utf-8 -*-
"""R7_six_criteria_scan CSV → Excel：Ticker 帶 TradingView 超連結；分頁：六項全中(淺紅 1/2/3 根) / 差一項 / 全部 / 說明。"""
import sys, glob, os, json
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

TV = "https://www.tradingview.com/chart/Q1c5VWwD/?symbol="
src = sys.argv[1] if len(sys.argv) > 1 else sorted(glob.glob(os.path.join(os.path.dirname(__file__), "R7_six_criteria_scan_r*.csv")))[-1]
d = pd.read_csv(src)
NC = 6
SRC_NOTE = os.environ.get("SRC_NOTE", "來源：三個 watchlist repo 的最新成品。")
BD = os.environ.get("BASE_DAY", "09-16")   # 六項條件的基準日（最新完整 OHLC 收盤）
SD = os.environ.get("SNAP_DAY", "09-17")   # 收盤快照日（只重算 ①④⑤⑥）
PANEL_N = os.environ.get("PANEL_N", "3,097")
VER = os.environ.get("VER", "r7")
# 版本說明（資料來源敘述、面板健康度）由呼叫端帶入，避免每版手改程式
DATA_NOTE = os.environ.get("DATA_NOTE", "")
PANEL_NOTE = os.environ.get("PANEL_NOTE", "")
FIX_NOTE = os.environ.get("FIX_NOTE", "")
SCOPE_NOTE = os.environ.get("SCOPE_NOTE", "")
COLS = [  # (csv 欄, Excel 標題, 格式)
    ("ticker", "Ticker", "link"), ("lists", "來源榜單", "txt"), ("date", "數據日", "txt"), ("close", "收盤", "num2"), ("score", "命中 /6", "int"),
    ("c1", "① 重心線向上", "tick"), ("grav", "重心線", "num2"), ("grav_slope", "重心斜率", "num3"), ("grav_shift5atr", "重心 5 根位移 (ATR)", "num2"),
    ("c2", "② 波動指數向上", "tick"), ("volidx", "波動指數", "num1"), ("volidx_slope", "波動指數斜率", "num2"), ("volidx_prev", "波動指數前值", "num1"),
    ("c3", "③ 波動指數 ≥75", "tick"), ("atr_pct", "ATR14 %", "num2"), ("sd60_pct", "60 日日波動 %", "num2"),
    ("c4", "④ 時鐘 6–10 點", "tick"), ("clock", "時鐘", "txt"), ("theta", "角度°", "num1"), ("cycle", "周期", "txt"), ("cyc_days", "周期第 N 日", "int"), ("cyc_prog", "周期進度 %", "int"),
    ("c5", "⑤ 淺紅第 1–3 根", "tick"), ("c5_cat", "淺紅第幾根", "txt"), ("hist", "Hist 今日", "num3"), ("hist_prev", "Hist 昨日", "num3"), ("hist_trough", "Hist 谷底", "num3"),
    ("c6", "⑥ 貼近重心線", "tick"), ("dist_grav_pct", "距重心 %", "num2"), ("dist_grav_atr", "距重心 (ATR)", "num2"),
    ("turnover20_m", "20 日均額 (百萬)", "num1"), ("bars", "歷史根數", "int"), ("hist_ok", "歷史足夠 ≥250", "tick"), ("data_warn", "資料警示", "txt"),
    ("vcp", "VCP 指數 (參考)", "num1"), ("vcp_grade", "VCP 等級", "txt"), ("struct", "結構", "txt"), ("mk", "Market Key", "txt"), ("lr_line", "最低阻力線", "num2"), ("dist_lr_pct", "距阻力線 %", "num2"),
]
if "ai_group" in d.columns:   # AI Sector watchlist 版：加小群組欄
    COLS[2:2] = [("ai_cat", "AI 大分類", "txt"), ("ai_group", "AI 小群組", "txt"), ("ai_rank", "小群組資金流排名", "int"), ("ai_flow5", "個股 5 日資金流向", "txt")]
    if "ai_chg5" in d.columns:
        COLS[6:6] = [("ai_chg5", "個股 5 日漲跌 %", "num2")]
d["hist_ok"] = d["bars"] >= 250
d["c5_cat"] = d.apply(lambda r: ("第 %d 根" % (int(r.c5_days) + 1)) if (pd.notna(r.c5_days) and r.still_light) else ("已中斷" if pd.notna(r.c5_days) else "—"), axis=1)
FMT = {"num1": "0.0", "num2": "0.00", "num3": "0.000", "int": "0"}
HEAD_FILL = PatternFill("solid", fgColor="EEF1ED"); OK = Font(color="16A34A", bold=True); NO = Font(color="DC2626", bold=True)
LINK = Font(color="0563C1", underline="single", bold=True); thin = Side(style="thin", color="E2E6E3")

def sheet(wb, name, df, note=None):
    ws = wb.create_sheet(name)
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = Font(italic=True, color="64748B"); r0 = 2
    for j, (_, h, _) in enumerate(COLS, 1):
        c = ws.cell(r0, j, h); c.font = Font(bold=True); c.fill = HEAD_FILL; c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, (_, r) in enumerate(df.iterrows(), r0 + 1):
        for j, (k, _, f) in enumerate(COLS, 1):
            v = r[k]
            if f == "link":
                c = ws.cell(i, j, str(v)); c.hyperlink = TV + str(v); c.font = LINK
            elif f == "tick":
                c = ws.cell(i, j, "✓" if bool(v) else "✗"); c.font = OK if bool(v) else NO; c.alignment = Alignment(horizontal="center")
            elif f in FMT:
                c = ws.cell(i, j, None if pd.isna(v) else float(v)); c.number_format = FMT[f]
            else:
                c = ws.cell(i, j, "" if pd.isna(v) else (str(v).replace("+", " · ") if k == "lists" else str(v)))
            c.border = Border(bottom=thin)
    ws.freeze_panes = ws.cell(r0 + 1, 2)
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(COLS))}{max(r0 + 1, r0 + len(df))}"
    for j, (k, h, f) in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(j)].width = 9 if f in ("tick", "num1", "num2", "num3", "int") else (34 if k == "ai_group" else max(8, min(26, len(h) * 1.6 + 2)))
    ws.row_dimensions[r0].height = 32
    return ws

wb = Workbook(); wb.remove(wb.active)
full = d[d.score == NC].sort_values(["c5_days", "volidx"], ascending=[True, False])
sheet(wb, "六項全中", full, "六項全中 %d 檔；按 ⑤ 淺紅第 1 / 2 / 3 根排列，類內依波動指數由高至低。Ticker 可點擊開 TradingView。" % len(full))
for k in (0, 1, 2):
    g = full[full.c5_days == k]
    sheet(wb, f"淺紅第 {k + 1} 根", g, f"六項全中 · 淺紅第 {k + 1} 根：{len(g)} 檔")
five = d[d.score == NC - 1].copy()
five["fail"] = five.apply(lambda r: [c for c in ("c1", "c2", "c3", "c4", "c5", "c6") if not r[c]][0], axis=1)
five = five.sort_values(["fail", "volidx"], ascending=[True, False])
sheet(wb, "差一項", five, "五中一 %d 檔；按缺少的條件分組（看 ✗ 在哪一欄）。" % len(five))
sheet(wb, "全部", d.sort_values(["score", "c5_days", "volidx"], ascending=[False, True, False]), "全部 %d 檔；用「命中 /6」欄篩選。" % len(d))
if "ai_group" in d.columns:
    sheet(wb, "按 AI 小群組", d.sort_values(["ai_rank", "score", "volidx"], ascending=[True, False, False]), "依 AI 小群組資金流排名 → 命中數 → 波動指數排列。")
# 全市場（可選）：把合併面板裡所有流動性達標的股票也跑一次，找出 watchlist 以外的命中
UNI = os.environ.get("UNIVERSE_CSV", "")
if UNI and os.path.exists(UNI):
    u = pd.read_csv(UNI)
    u["hist_ok"] = u["bars"] >= 250
    u["c5_cat"] = u.apply(lambda r: ("第 %d 根" % (int(r.c5_days) + 1)) if (pd.notna(r.c5_days) and r.still_light) else ("已中斷" if pd.notna(r.c5_days) else "—"), axis=1)
    inwl = set(d.ticker)
    uf = u[u.score == NC].copy()
    uf["lists"] = uf.ticker.map(lambda t: "榜內" if t in inwl else "榜外")
    uf = uf.sort_values(["c5_days", "volidx"], ascending=[True, False])
    sheet(wb, "加掃 六項全中", uf, f"三鏡像併集 {PANEL_N} 檔中、收盤 ≥ $2、20 日 (收盤×量) 平均 ≥ 300 萬美元、歷史 ≥ 80 根、最後一根在 {BD}、同證券不同寫法只留一個的 {len(u):,} 檔，六項全中 %d 檔（其中 %d 檔不在五份來源成品內）。" % (len(uf), (uf.lists == "榜外").sum()) + os.environ.get("MKT_NOTE", ""))

# chat 1–3 命中但未過六項（可選）：三個 session 各自用自己那套準則選中、但六項條件未全中的名字
CH = os.environ.get("CHATHITS_CSV", "")
if CH and os.path.exists(CH):
    ch = pd.read_csv(CH)
    CHCOLS = [("ticker", "Ticker", "link"), ("name", "公司", "txt"),
              ("n_hit", "命中系統數", "int"), ("hits", "命中哪幾套", "txt"), ("detail", "命中內容（來源原話）", "txt"),
              ("avoid", "反向標註", "txt"), ("catalyst", "催化", "txt"),
              ("score", "六項命中 /6", "int"), ("miss", "欠缺條件", "txt"),
              ("close", f"{BD} 收盤", "num2"), ("volidx", "波動指數", "num1"), ("volidx_slope", "指數斜率", "num2"),
              ("grav_slope", "重心斜率", "num3"), ("clock", "時鐘", "txt"), ("c5_cat", "淺紅第幾根", "txt"),
              ("dist_grav_pct", "距重心 %", "num2"),
              ("comb_grade", "chat2 VCP/Wein/Pre", "txt"), ("comb_up", "chat2 上升分數", "num1"), ("comb_sure", "chat2 確定性", "num1"),
              ("ss_name", "chat3 子板塊", "txt"), ("ss_rank", "子板塊名次", "int"), ("ss_score", "子板塊 5 日分", "num1"),
              ("ai_group", "chat3 AI 小群組", "txt"), ("ai_rank", "AI 群組名次", "int"), ("ai_tf5", "AI 個股 5 日強度", "num3"),
              ("mp_rank", "chat1 R21 排名", "int"), ("mp_win", "R21 動能時間框", "txt"), ("mp_score", "R21 爆發潛力", "num1"),
              ("mp_flag", "R21 審視標記", "txt"), ("rh", "chat3 加息 R2", "txt"),
              ("four_17", f"{SD} ①④⑤⑥", "tick"), ("close_17", f"{SD} 收盤", "num2"), ("chg1d_pct", f"{BD}→{SD} %", "num2"),
              ("turnover20_m", "20 日均額 (百萬)", "num1"), ("bars", "歷史根數", "int"), ("hist_ok", "歷史足夠 ≥250", "tick"),
              ("vcp", "VCP 指數 (參考)", "num1"), ("struct", "結構", "txt"), ("lists", "來源榜單", "txt")]

    if not (os.environ.get("UPDATE_CSV") and os.path.exists(os.environ.get("UPDATE_CSV"))):
        CHCOLS = [c for c in CHCOLS if c[0] not in ("four_17", "close_17", "chg1d_pct")]

    def chsheet(name, df, note):
        w = wb.create_sheet(name)
        w.cell(1, 1, note).font = Font(italic=True, color="64748B")
        for j, (_, h, _) in enumerate(CHCOLS, 1):
            c = w.cell(2, j, h); c.font = Font(bold=True); c.fill = HEAD_FILL
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i, (_, r) in enumerate(df.iterrows(), 3):
            for j, (k, _, f) in enumerate(CHCOLS, 1):
                v = r[k]
                if f == "link":
                    c = w.cell(i, j, str(v)); c.hyperlink = TV + str(v); c.font = LINK
                elif f == "tick":
                    c = w.cell(i, j, "" if pd.isna(v) else ("✓" if bool(v) else "✗"))
                    c.font = OK if bool(v) else NO; c.alignment = Alignment(horizontal="center")
                elif f in FMT:
                    c = w.cell(i, j, None if pd.isna(v) else float(v)); c.number_format = FMT[f]
                else:
                    c = w.cell(i, j, "" if pd.isna(v) else (str(v).replace("+", " · ") if k == "lists" else str(v)))
                c.border = Border(bottom=thin)
        w.freeze_panes = w.cell(3, 2)
        w.auto_filter.ref = f"A2:{get_column_letter(len(CHCOLS))}{max(3, 2 + len(df))}"
        for j, (k, h, f) in enumerate(CHCOLS, 1):
            w.column_dimensions[get_column_letter(j)].width = (
                58 if k == "detail" else 30 if k in ("hits", "ai_group", "catalyst", "mp_flag", "lists") else
                20 if k in ("name", "ss_name", "comb_grade") else
                9 if f in ("tick", "num1", "num2", "num3", "int") else 12)
        w.row_dimensions[2].height = 32
        return w

    chsheet("Chat1-3 命中 未過六項", ch,
            "三個 session 各自用自己那套準則選中、但 R7 六項條件未全中的 %d 檔。命中規則（一律採用來源成品自己的通過標記）："
            "chat1 動能回調 R21 = 在總表（1/2/3/6 個月動能至少一個排前 10%% 且回到上升中的 20MA；差一項名單不算）；"
            "chat2 Combined R22 = 線上 ≥1 且三榜有頂級（VCP A/B、Weinstein 2A、Pre-breakout A）；"
            "chat3 SubSector R12 = 所屬子板塊 5 日分 ≥70 且斜率 >0；"
            "chat3 AI Sector R13 = 所屬小群組 5 日分 ≥70 且個股 5 日強度 >0；chat3 加息 R2 = 受惠名單。"
            "按「六項命中 /6」→「命中系統數」→ 波動指數排序。" % len(ch))
    for k in (3, 2, 1):
        g = ch[ch.n_hit == k]
        if len(g):
            chsheet(f"Chat 命中 {k} 套", g, f"同時命中 {k} 套 chat 準則、但六項未全中：{len(g)} 檔。")

# 收盤快照更新（可選）：快照沒有盤中高低，②③（波動指數）不可信，只看 ①④⑤⑥
UPD = os.environ.get("UPDATE_CSV", "")
if UPD and os.path.exists(UPD):
    m = pd.read_csv(UPD)
    UCOLS = [("ticker", "Ticker", "link"), ("lists", "來源榜單", "txt"), ("close_16", f"{BD} 收盤", "num2"),
             ("close_17", f"{SD} 收盤", "num2"), ("chg1d_pct", f"{BD}→{SD} %", "num2"),
             ("c1_17", "① 重心線向上", "tick"), ("c4_17", "④ 時鐘 6–10 點", "tick"), ("c5_17", "⑤ 淺紅 1–3 根", "tick"),
             ("c6_17", "⑥ 貼近重心線", "tick"), ("clock_17", f"{SD} 時鐘", "txt"), ("dist_grav_pct", "距重心 %", "num2"),
             ("volidx", f"波動指數 ({BD})", "num1"), ("c2", f"② {BD} 指數向上", "tick"), ("c3", f"③ {BD} ≥75", "tick")]

    def usheet(name, df, note):
        ws2 = wb.create_sheet(name)
        ws2.cell(1, 1, note).font = Font(italic=True, color="64748B")
        for j, (_, h, _) in enumerate(UCOLS, 1):
            c = ws2.cell(2, j, h); c.font = Font(bold=True); c.fill = HEAD_FILL
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i, (_, r) in enumerate(df.iterrows(), 3):
            for j, (k, _, f) in enumerate(UCOLS, 1):
                v = r[k]
                if f == "link":
                    c = ws2.cell(i, j, str(v)); c.hyperlink = TV + str(v); c.font = LINK
                elif f == "tick":
                    c = ws2.cell(i, j, "✓" if bool(v) else "✗"); c.font = OK if bool(v) else NO
                    c.alignment = Alignment(horizontal="center")
                elif f in FMT:
                    c = ws2.cell(i, j, None if pd.isna(v) else float(v)); c.number_format = FMT[f]
                else:
                    c = ws2.cell(i, j, "" if pd.isna(v) else (str(v).replace("+", " · ") if k == "lists" else str(v)))
                c.border = Border(bottom=thin)
        ws2.freeze_panes = ws2.cell(3, 2)
        for j, (k, h, f) in enumerate(UCOLS, 1):
            ws2.column_dimensions[get_column_letter(j)].width = 10 if f in ("tick", "num1", "num2") else 22
        ws2.row_dimensions[2].height = 30

    still = m[(m.score == NC) & m.four_17]
    lost = m[(m.score == NC) & (~m.four_17)]
    fresh = m[(m.score != NC) & m.four_17 & m.c2 & m.c3].sort_values("volidx", ascending=False)
    usheet(f"{SD} 仍成立", still, f"{BD} 六項全中、且在 {SD} 收盤下 ①④⑤⑥ 仍全部成立：{len(still)} / {len(still) + len(lost)} 檔。")
    usheet(f"{SD} 已失效", lost, f"{BD} 六項全中、但 {SD} 收盤已有一項不成立（多數是 ⑤ 淺紅中斷）：{len(lost)} 檔。")
    usheet(f"{SD} 新符合", fresh, f"{SD} 收盤下 ①④⑤⑥ 成立、且 {BD} 的 ②③（波動指數）也通過的候補：{len(fresh)} 檔。快照無盤中高低，②③ 用 {BD} 值。")

# 與上版對照（可選）：上一版六項全中與本版六項全中的延續／新進／退出
PREV = os.environ.get("PREV_CSV", "")
PV = os.environ.get("PREV_VER", "上版"); PD_ = os.environ.get("PREV_DAY", "")
if PREV and os.path.exists(PREV):
    pv = pd.read_csv(PREV).set_index("ticker")
    cur = d.set_index("ticker")
    CNM = {"c1": "①", "c2": "②", "c3": "③", "c4": "④", "c5": "⑤", "c6": "⑥"}
    rows = []
    for t in list(full.ticker) + [t for t in pv.index[pv.score == NC] if t not in set(full.ticker)]:
        p6 = t in pv.index and int(pv.loc[t, "score"]) == NC
        c = cur.loc[t] if t in cur.index else None
        st = "延續" if (p6 and c is not None and int(c.score) == NC) else ("新進" if not p6 else ("退出" if c is not None else "已不在來源榜單"))
        rows.append(dict(ticker=t, status=st, lists=(c["lists"] if c is not None else pv.loc[t, "lists"]),
                         prev_score=int(pv.loc[t, "score"]) if t in pv.index else None,
                         score=int(c.score) if c is not None else None,
                         miss="".join(CNM[k] for k in CNM if c is not None and not bool(c[k])),
                         close=c.close if c is not None else None, volidx=c.volidx if c is not None else None,
                         clock=c.clock if c is not None else "", c5_cat=c.c5_cat if c is not None else ""))
    cmp_ = pd.DataFrame(rows)
    cmp_["_o"] = cmp_.status.map({"延續": 0, "新進": 1, "退出": 2, "已不在來源榜單": 3})
    cmp_ = cmp_.sort_values(["_o", "volidx"], ascending=[True, False]).drop(columns="_o")
    PCOLS = [("ticker", "Ticker", "link"), ("status", "狀態", "txt"), ("lists", "來源榜單", "txt"),
             ("prev_score", f"{PV} 命中 /6（{PD_}）", "int"), ("score", f"本版命中 /6（{BD}）", "int"), ("miss", "本版欠缺", "txt"),
             ("close", f"{BD} 收盤", "num2"), ("volidx", "波動指數", "num1"), ("clock", "時鐘", "txt"), ("c5_cat", "淺紅第幾根", "txt")]
    wp = wb.create_sheet(f"與 {PV} 對照")
    k_ = cmp_.status.value_counts().to_dict()
    wp.cell(1, 1, f"{PV}（{PD_} 基準）六項全中 {int((pv.score == NC).sum())} 檔 vs 本版（{BD} 基準）{len(full)} 檔："
                  f"延續 {k_.get('延續', 0)}、新進 {k_.get('新進', 0)}、退出 {k_.get('退出', 0) + k_.get('已不在來源榜單', 0)}"
                  + (f"（其中 {k_['已不在來源榜單']} 檔已不在來源榜單）" if k_.get('已不在來源榜單') else "") + "。").font = Font(italic=True, color="64748B")
    for j, (_, h, _) in enumerate(PCOLS, 1):
        c = wp.cell(2, j, h); c.font = Font(bold=True); c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, (_, r) in enumerate(cmp_.iterrows(), 3):
        for j, (k, _, f) in enumerate(PCOLS, 1):
            v = r[k]
            if f == "link":
                c = wp.cell(i, j, str(v)); c.hyperlink = TV + str(v); c.font = LINK
            elif f in FMT:
                c = wp.cell(i, j, None if pd.isna(v) else float(v)); c.number_format = FMT[f]
            else:
                c = wp.cell(i, j, "" if pd.isna(v) else (str(v).replace("+", " · ") if k == "lists" else str(v)))
            c.border = Border(bottom=thin)
    wp.freeze_panes = wp.cell(3, 2)
    for j, (k, h, f) in enumerate(PCOLS, 1):
        wp.column_dimensions[get_column_letter(j)].width = 30 if k == "lists" else (12 if f in ("num1", "num2", "int") else 14)
    wp.row_dimensions[2].height = 32

# 來源分頁（可選）：三個 session 的最新成品出處
SRC_ROWS = os.environ.get("SRC_ROWS", "")
if SRC_ROWS and os.path.exists(SRC_ROWS):
    srcs = pd.read_csv(SRC_ROWS)
    wsx = wb.create_sheet("來源")
    hdr = list(srcs.columns)
    for j, h in enumerate(hdr, 1):
        c = wsx.cell(1, j, h); c.font = Font(bold=True); c.fill = HEAD_FILL
    for i, (_, r) in enumerate(srcs.iterrows(), 2):
        for j, h in enumerate(hdr, 1):
            v = r[h]
            if pd.isna(v):
                cell = ""
            else:
                cell = int(v) if (h.endswith("檔數") and str(v).strip().isdigit()) else str(v)
            wsx.cell(i, j, cell)
    for j, h in enumerate(hdr, 1):
        wsx.column_dimensions[get_column_letter(j)].width = 14 if h.endswith(("檔數", "基準日")) else 46

CH_NOTE = UNI_NOTE = ""
if CH and os.path.exists(CH):
    _bj = CH[:-4] + "_both.json"          # chat_hits.py 另寫的「同時六項全中」名單
    _bjd = json.load(open(_bj)) if os.path.exists(_bj) else None
    _ch_all = set(_bjd["both"]) if _bjd else None
    _bav = set(_bjd.get("both_avoid", [])) if _bjd else set()
    _sav = [t for t in (_bjd.get("six_avoid", []) if _bjd else []) if t not in (_ch_all or set())]
    _ov = [t for t in full.ticker if _ch_all is not None and t in _ch_all]
    CH_NOTE = ("「Chat1-3 命中 未過六項」分頁：三個 session 各自的成品用自己那套準則選中、但本掃描六項未全中的名字，另按命中 1 / 2 / 3 套拆成子分頁。命中規則寫在該分頁第一行。"
               + (f"本版六項全中 {len(full)} 檔裡有 {len(_ov)} 檔同時是 chat 命中（{'、'.join(t + ('（加息清單「迴避」）' if t in _bav else '') for t in _ov)}），其餘 {len(full) - len(_ov)} 檔是本掃描獨有"
          + (f"（其中 {'、'.join(_sav)} 在加息清單「迴避」）" if _sav else "") + "。" if _ch_all is not None else ""))
if UNI and os.path.exists(UNI):
    UNI_NOTE = f"「加掃 六項全中」分頁：合併面板中收盤 ≥ $2、20 日 (收盤×量) 平均 ≥ 300 萬美元、歷史 ≥ 80 根、最後一根在 {BD}、同證券不同寫法只留一個的 {len(u):,} 檔全掃一次的結果（六項全中 {len(uf)} 檔，其中 {(uf.lists == '榜外').sum()} 檔不在五份來源成品內），含 watchlist 以外的名字。"

ws = wb.create_sheet("說明")
for i, t in enumerate([t for t in [
    f"R7 A–H 六項條件掃描 {VER} — 六項條件的基準是 2026-{BD} 官方收盤（完整 OHLC，三個 repo 的 Yahoo 鏡像合併 {PANEL_N} 檔）。"
    + (f"另有三個 {SD} 分頁，用 10MA repo 的 Nasdaq 收盤快照更新 ①④⑤⑥；快照沒有盤中高低，②③（波動指數）不能重算，沿用 {BD} 值。" if (UPD and os.path.exists(UPD)) else ""),
    DATA_NOTE,
    "① 重心線向上：r7e_gravity（滾動 VWAP 30，hlc3）最新一根斜率 > 0",
    "② 波動指數向上：r7h_volidx（0–100，高 = 平靜）最新一根斜率 > 0",
    "③ 波動指數 ≥ 75",
    "④ 時鐘 6–10 點：r7b_macd_clock 指針 180°–300°",
    "⑤ 淺紅第 1–3 根：MACD 柱狀圖 < 0 且回升，連續 1–3 根、未中斷",
    "⑥ 貼近重心線：收盤距重心線 ≤ 1.0 × ATR14 或 ≤ 3%（% 以收盤價為分母，與「距重心 %」欄相同）",
    "VCP / 結構 / 最低阻力線 只作參考，不計分。MA20 與 EMA21 已刪除。",
    "Ticker 欄為 TradingView 圖表超連結（Q1c5VWwD 版面）。",
    "來源榜單欄會標明名字的身分：MP_R21（chat 1 動能回調總表）· MP_R21_差一項（形態只差一項，只掃描）· RateHike_R2_受惠 / _迴避 / _索引（chat 3 SubSector session 09-18 的加息與地緣政治 3 日清單）。六項全中裡若出現「差一項」或「迴避」標籤，代表原榜單本身沒有選中或不推薦，請自行判斷。",
    PANEL_NOTE,
    FIX_NOTE,
    SCOPE_NOTE or f"{len(d)} 檔全部掃到，0 檔缺資料。",
    CH_NOTE,
    UNI_NOTE,
    "① 用的是「重心線 1 根斜率 > 0」（六項條件的原文）。R7-E 自己把線塗綠的條件較嚴：5 根位移 ≥ 0.8 ATR；表中已附「重心 5 根位移 (ATR)」欄，可自行加嚴。",
    "③ 門檻用 75（六項條件原文）。R7-H 圖上那條綠色分隔線預設在 80，所以有些通過 ③ 的名字在圖上仍在線下。",
    "④ 時鐘：MACD 柱狀圖在 Pine 的前 33 根是 na（EMA 要等 SMA 種子），本掃描已照做，並丟掉含第一根有效柱的那一段（它在視窗起點之前就開始，長度被截斷）；其後每一段都是完整週期，取最近 8 段平均，與 TradingView 載入長歷史時的結果一致（r7 以前多丟了另一方向的第一段完整週期；r8 起已修正）。歷史根數 < 150 的名字，時鐘與進度仍不夠可靠，請以「歷史根數」欄判斷。",
    "代號對齊：鏡像用 BRK/A、BRK/B、BF/B 等斜線寫法，watchlist 用 BRK-A、BF.B；掃描會自動試點/槓/斜線並取資料最長者（本次 " + os.environ.get("REMAP_NOTE", "3 檔重對應") + "）。",
    SRC_NOTE,
] if t], 1):
    ws.cell(i, 1, t).font = Font(bold=(i == 1))
ws.column_dimensions["A"].width = 110
out = sys.argv[2] if len(sys.argv) > 2 else src[:-4] + ".xlsx"
wb.save(out); print(out)
